"""
Shared helpers for the Document Manager and DQAS Certificate apps —
QR/image rendering, printable sheets, auth, and small date utilities.
Both apps import this module rather than duplicating this logic.
"""

import streamlit as st
import sqlite3
import qrcode
import hashlib
import requests
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font as XLFont
from openpyxl.utils import get_column_letter
import io
import os
import calendar
from datetime import datetime
import fitz  # PyMuPDF — used only to rasterize PDF certificates for stamping

# sha256 of "Dny@nesh57"
_CREDENTIALS = {
    "dnyanesh.nikam@sjcpl.in": hashlib.sha256(b"Dny@nesh57").hexdigest(),
}

_HERE = os.path.dirname(os.path.abspath(__file__))


# ----------------------------------------------------------------------
# Database layer (generic — each app passes its own db_path)
# ----------------------------------------------------------------------

def get_conn(db_path):
    conn = sqlite3.connect(db_path, check_same_thread=False, timeout=30)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA busy_timeout = 30000")
    return conn


def get_setting(db_path, key, default=""):
    conn = get_conn(db_path)
    row = conn.execute("SELECT value FROM settings WHERE key=?", (key,)).fetchone()
    conn.close()
    return row[0] if row else default


def set_setting(db_path, key, value):
    conn = get_conn(db_path)
    conn.execute(
        "INSERT INTO settings (key, value) VALUES (?, ?) "
        "ON CONFLICT(key) DO UPDATE SET value=excluded.value",
        (key, value),
    )
    conn.commit()
    conn.close()


# --- Generic named lookup tables (projects/categories/departments/sites/...) ---

def get_lookup(db_path, table):
    conn = get_conn(db_path)
    rows = conn.execute(f"SELECT name FROM {table} ORDER BY name").fetchall()
    conn.close()
    return [r[0] for r in rows]


def add_lookup(db_path, table, name):
    name = name.strip()
    if not name:
        return False
    conn = get_conn(db_path)
    try:
        conn.execute(f"INSERT INTO {table} (name) VALUES (?)", (name,))
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False
    finally:
        conn.close()


def delete_lookup(db_path, table, name):
    conn = get_conn(db_path)
    conn.execute(f"DELETE FROM {table} WHERE name=?", (name,))
    conn.commit()
    conn.close()


# ----------------------------------------------------------------------
# Auth
# ----------------------------------------------------------------------

def check_password(email, password):
    h = _CREDENTIALS.get(email.strip().lower())
    return h is not None and h == hashlib.sha256(password.encode()).hexdigest()


def show_login(app_title, app_caption, page_icon="🔐"):
    st.set_page_config(page_title=f"{app_title} — Login", page_icon=page_icon, layout="centered")
    st.title("🔐 Sign in")
    st.caption(app_caption)
    with st.form("login_form"):
        email = st.text_input("Email")
        password = st.text_input("Password", type="password")
        submitted = st.form_submit_button("Sign in", type="primary")
        if submitted:
            if check_password(email, password):
                st.session_state["authenticated"] = True
                st.rerun()
            else:
                st.error("Invalid email or password.")


# ----------------------------------------------------------------------
# QR / image helpers
# ----------------------------------------------------------------------

def _load_fonts(size_bold, size_regular):
    candidates_bold = [
        os.path.join(_HERE, "Roboto-Bold.ttf"),
        "arialbd.ttf", r"C:\Windows\Fonts\arialbd.ttf",
        r"C:\Windows\Fonts\calibrib.ttf",
    ]
    candidates_regular = [
        os.path.join(_HERE, "Roboto-Regular.ttf"),
        "arial.ttf", r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibri.ttf",
    ]
    def _try(paths, size):
        for p in paths:
            try:
                return ImageFont.truetype(p, size)
            except Exception:
                pass
        return ImageFont.load_default()
    return _try(candidates_bold, size_bold), _try(candidates_regular, size_regular)


def _wrap_lines(draw, text, font, max_width):
    """Word-wrap text to fit max_width, hard-breaking any single word that
    alone exceeds max_width so long names can never overflow a fixed area."""
    if not text:
        return [""]
    words = text.split()
    lines, current = [], ""
    for word in words:
        while draw.textlength(word, font=font) > max_width and len(word) > 1:
            cut = len(word)
            while cut > 1 and draw.textlength(word[:cut], font=font) > max_width:
                cut -= 1
            if current:
                lines.append(current)
                current = ""
            lines.append(word[:cut])
            word = word[cut:]
        trial = f"{current} {word}".strip()
        if not current or draw.textlength(trial, font=font) <= max_width:
            current = trial
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def _line_height(font):
    bbox = font.getbbox("Agjpqy")
    return (bbox[3] - bbox[1]) + 8


def _make_qr_pixels(data, target_px):
    probe = qrcode.QRCode(
        version=None, error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=1, border=4,
    )
    probe.add_data(data)
    probe.make(fit=True)
    modules = probe.modules_count + 8
    box_size = max(1, target_px // modules)
    qr = qrcode.QRCode(
        version=None, error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=box_size, border=4,
    )
    qr.add_data(data)
    qr.make(fit=True)
    return qr.make_image(fill_color="black", back_color="white").convert("RGB")


def make_qr_image(data, box_size=8, border=4):
    qr = qrcode.QRCode(
        version=None, error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=box_size, border=border,
    )
    qr.add_data(data)
    qr.make(fit=True)
    return qr.make_image(fill_color="black", back_color="white").convert("RGB")


def make_labeled_qr(item_id, name, data, target_px=1200):
    qr_img = _make_qr_pixels(data, target_px)
    w = qr_img.width
    font_size_id = max(48, w // 9)
    font_size_name = max(36, w // 12)
    font_id, font_name = _load_fonts(font_size_id, font_size_name)

    metrics = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    max_text_w = w - 40
    id_lines = _wrap_lines(metrics, item_id, font_id, max_text_w)
    name_lines = _wrap_lines(metrics, name, font_name, max_text_w)
    id_line_h = _line_height(font_id)
    name_line_h = _line_height(font_name)
    gap = 16
    text_h = len(id_lines) * id_line_h + gap + len(name_lines) * name_line_h
    label_h = text_h + 60

    canvas = Image.new("RGB", (w, w + label_h), "white")
    canvas.paste(qr_img, (0, 0))
    draw = ImageDraw.Draw(canvas)

    ty = w + 20
    for line in id_lines:
        lw = draw.textlength(line, font=font_id)
        draw.text(((w - lw) / 2, ty), line, fill="black", font=font_id)
        ty += id_line_h
    ty += gap
    for line in name_lines:
        lw = draw.textlength(line, font=font_name)
        draw.text(((w - lw) / 2, ty), line, fill="#333333", font=font_name)
        ty += name_line_h
    return canvas


def img_to_bytes(img, dpi=300):
    buf = io.BytesIO()
    img.save(buf, format="PNG", dpi=(dpi, dpi))
    return buf.getvalue()


def images_to_pdf_bytes(images):
    """Combine a list of PIL Images into a single multi-page PDF (bytes)."""
    buf = io.BytesIO()
    imgs = [im.convert("RGB") for im in images]
    imgs[0].save(buf, format="PDF", save_all=True, append_images=imgs[1:])
    return buf.getvalue()


def build_qr_excel(rows, columns, qr_col_header="QR Code"):
    """Builds an .xlsx (as bytes): one row per item, with the given text
    columns plus a QR code image embedded in the last column — for handing
    to a print vendor so they can match each QR to the right certificate.

    rows: list of dicts, each with a 'qr_image' PIL Image plus whatever
    keys `columns` references.
    columns: ordered list of (header, dict_key) tuples for the text columns.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "QR Codes"

    headers = [h for h, _ in columns] + [qr_col_header]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        ws.cell(row=1, column=col_idx).font = XLFont(bold=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = 18
    qr_col = len(columns) + 1

    image_buffers = []  # keep alive until wb.save()
    for row_idx, item in enumerate(rows, start=2):
        for col_idx, (_, key) in enumerate(columns, start=1):
            ws.cell(row=row_idx, column=col_idx, value=item.get(key, "") or "")
        ws.row_dimensions[row_idx].height = 90

        qr_img = item["qr_image"].resize((110, 110))
        buf = io.BytesIO()
        qr_img.save(buf, format="PNG")
        buf.seek(0)
        image_buffers.append(buf)
        xl_img = XLImage(buf)
        ws.add_image(xl_img, f"{get_column_letter(qr_col)}{row_idx}")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_print_sheet(items, title=None):
    """Two-column print table — QR Code | Title.
    items: list of (item_id, name, data) tuples.
    If title is given, it's rendered as a plain heading above the table
    (e.g. a category name) — not part of the table itself.
    Style: white bg, #00AEDA accent header, black text, black borders,
    Roboto font, plain white rows throughout (no banding).
    Resolution: 300 DPI equivalent (high quality for print).
    Row height grows to fit wrapped names, so text can never overlap
    the QR column regardless of how long a name is.
    """
    ACCENT = (0, 174, 218)    # #00AEDA
    WHITE = (255, 255, 255)
    BLACK = (0, 0, 0)

    # All sizes are in pixels at ~300 DPI
    qr_px = 900            # QR image size per cell
    pad = 50                # inner cell padding
    line = 6                # border thickness

    col1_w = qr_px + pad * 2
    col2_w = 1300
    hdr_h = 160
    title_h = 140 if title else 0

    title_font_sz = 80
    hdr_font_sz = 90
    id_font_sz = 64
    name_font_sz = 52
    text_gap = 20

    font_title, _ = _load_fonts(title_font_sz, title_font_sz)
    font_hdr, _ = _load_fonts(hdr_font_sz, hdr_font_sz)
    font_id, font_name = _load_fonts(id_font_sz, name_font_sz)

    total_w = col1_w + line + col2_w
    x_div = col1_w
    max_text_w = col2_w - pad * 2
    id_line_h = _line_height(font_id)
    name_line_h = _line_height(font_name)

    metrics = ImageDraw.Draw(Image.new("RGB", (1, 1)))

    # ── Pass 1: compute wrapped lines + row height for every item ─────────
    rows = []
    for item_id, name, data in items:
        id_lines = _wrap_lines(metrics, item_id, font_id, max_text_w)
        name_lines = _wrap_lines(metrics, name, font_name, max_text_w)
        text_h = len(id_lines) * id_line_h + text_gap + len(name_lines) * name_line_h
        row_h = max(qr_px + pad * 2, text_h + pad * 2)
        rows.append({
            "item_id": item_id, "data": data,
            "id_lines": id_lines, "name_lines": name_lines, "row_h": row_h,
        })

    total_h = title_h + (line if title else 0) + hdr_h + line
    total_h += sum(r["row_h"] + line for r in rows)

    sheet = Image.new("RGB", (total_w, total_h), WHITE)
    draw = ImageDraw.Draw(sheet)
    y = 0

    # ── Title (above the table, plain white background) ─────────────────
    if title:
        tw = draw.textlength(title, font=font_title)
        draw.text(((total_w - tw) / 2, (title_h - title_font_sz) / 2), title, fill=BLACK, font=font_title)
        y += title_h
        draw.rectangle([(0, y), (total_w, y + line)], fill=BLACK)
        y += line

    # ── Header row ─────────────────────────────────────────────────────
    draw.rectangle([(0, y), (total_w, y + hdr_h)], fill=ACCENT)
    for text, x_start, col_w in [
        ("QR Code", 0, col1_w),
        ("Title", x_div + line, col2_w),
    ]:
        tw = draw.textlength(text, font=font_hdr)
        tx = x_start + (col_w - tw) / 2
        ty = y + (hdr_h - hdr_font_sz) / 2
        draw.text((tx, ty), text, fill=WHITE, font=font_hdr)
    draw.rectangle([(x_div, y), (x_div + line, y + hdr_h)], fill=WHITE)
    y += hdr_h
    draw.rectangle([(0, y), (total_w, y + line)], fill=BLACK)
    y += line

    # ── Data rows (plain white, no banding) ──────────────────────────────
    for r in rows:
        row_h = r["row_h"]
        y0, y1 = y, y + row_h

        qr_img = _make_qr_pixels(r["data"], qr_px)
        qx = (col1_w - qr_img.width) // 2
        qy = y0 + (row_h - qr_img.height) // 2
        sheet.paste(qr_img, (qx, qy))

        text_h = len(r["id_lines"]) * id_line_h + text_gap + len(r["name_lines"]) * name_line_h
        ty = y0 + (row_h - text_h) // 2
        for l in r["id_lines"]:
            lw = draw.textlength(l, font=font_id)
            draw.text((x_div + line + (col2_w - lw) / 2, ty), l, fill=BLACK, font=font_id)
            ty += id_line_h
        ty += text_gap
        for l in r["name_lines"]:
            lw = draw.textlength(l, font=font_name)
            draw.text((x_div + line + (col2_w - lw) / 2, ty), l, fill=BLACK, font=font_name)
            ty += name_line_h

        draw.rectangle([(x_div, y0), (x_div + line, y1)], fill=BLACK)
        draw.rectangle([(0, y1), (total_w, y1 + line)], fill=BLACK)
        y = y1 + line

    # ── Outer border ───────────────────────────────────────────────────
    draw.rectangle([(0, 0), (total_w - 1, total_h - 1)], outline=BLACK, width=line)

    return sheet


def build_category_sheets(items, per_page=4):
    """items: list of (item_id, name, category, data).
    Returns an ordered {label: PIL Image} dict — one print sheet per page.
    Each category (alphabetical, 'Uncategorized' last) is split into pages
    of at most `per_page` items so rows stay a comfortable, readable
    size. A category with more than `per_page` items becomes multiple
    sheets titled "Category (Page X of Y)"; one that fits on a single page
    keeps a plain "Category" title."""
    groups = {}
    for item_id, name, category, data in items:
        cat = (category or "").strip() or "Uncategorized"
        groups.setdefault(cat, []).append((item_id, name, data))

    def _sort_key(cat):
        return (cat == "Uncategorized", cat.lower())

    sheets = {}
    for cat in sorted(groups.keys(), key=_sort_key):
        entries = groups[cat]
        pages = [entries[i:i + per_page] for i in range(0, len(entries), per_page)]
        total_pages = len(pages)
        for page_idx, page_entries in enumerate(pages, start=1):
            title = cat if total_pages == 1 else f"{cat} (Page {page_idx} of {total_pages})"
            sheets[title] = build_print_sheet(page_entries, title=title)
    return sheets


# ----------------------------------------------------------------------
# Date helpers (used by the DQAS certificate app)
# ----------------------------------------------------------------------

def add_months(d, months):
    """Add a whole number of months to a date, clamping the day to the
    last valid day of the target month (e.g. 31 Jan + 1 month = 28/29 Feb)."""
    month_index = d.month - 1 + months
    year = d.year + month_index // 12
    month = month_index % 12 + 1
    day = min(d.day, calendar.monthrange(year, month)[1])
    return d.replace(year=year, month=month, day=day)


def compute_expiry_date(issue_date_str, validity_months):
    issue = datetime.strptime(issue_date_str, "%Y-%m-%d").date()
    return add_months(issue, validity_months)


# ----------------------------------------------------------------------
# Certificate stamping — overlay a small QR + ID stamp onto a certificate
# file at a fixed, manually-calibrated position.
# ----------------------------------------------------------------------

def make_qr_id_stamp(cert_id, data, target_w):
    """A compact QR code with the certificate ID printed below it —
    the unit that gets stamped onto a certificate file."""
    qr_img = _make_qr_pixels(data, target_w)
    w = qr_img.width
    font_size = max(28, w // 10)
    font_id, _ = _load_fonts(font_size, font_size)

    metrics = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    max_text_w = w - 20
    id_lines = _wrap_lines(metrics, cert_id, font_id, max_text_w)
    line_h = _line_height(font_id)
    label_h = len(id_lines) * line_h + 20

    canvas = Image.new("RGB", (w, w + label_h), "white")
    canvas.paste(qr_img, (0, 0))
    draw = ImageDraw.Draw(canvas)

    ty = w + 10
    for line in id_lines:
        lw = draw.textlength(line, font=font_id)
        draw.text(((w - lw) / 2, ty), line, fill="black", font=font_id)
        ty += line_h
    return canvas


def load_certificate_image(source, dpi=200):
    """Load a certificate as a PIL Image, regardless of whether it's an
    image file or a PDF (first page only). `source` can be a filesystem
    path (str) or a file-like object (e.g. from st.file_uploader)."""
    if hasattr(source, "read"):
        name = getattr(source, "name", "")
        data = source.read()
        ext = os.path.splitext(name)[1].lower()
    else:
        ext = os.path.splitext(source)[1].lower()
        with open(source, "rb") as f:
            data = f.read()

    if ext == ".pdf":
        doc = fitz.open(stream=data, filetype="pdf")
        page = doc[0]
        zoom = dpi / 72
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        doc.close()
        return img

    return Image.open(io.BytesIO(data)).convert("RGB")


def stamp_certificate(base_img, cert_id, data, x_pct, y_pct, size_pct):
    """Paste a QR+ID stamp onto a copy of base_img. x_pct/y_pct position the
    stamp's center as a percentage of the certificate's width/height;
    size_pct sets the stamp's width as a percentage of the certificate's
    width. The stamp is clamped to stay fully within the certificate."""
    base = base_img.convert("RGB").copy()
    target_w = max(40, int(base.width * size_pct / 100))
    stamp = make_qr_id_stamp(cert_id, data, target_w)

    px = int(base.width * x_pct / 100) - stamp.width // 2
    py = int(base.height * y_pct / 100) - stamp.height // 2
    px = max(0, min(px, max(0, base.width - stamp.width)))
    py = max(0, min(py, max(0, base.height - stamp.height)))

    base.paste(stamp, (px, py))
    return base


def load_overlay_image(source):
    """Load a transparency-preserving overlay image (e.g. a signature PNG).
    `source` can be a filesystem path or a file-like object."""
    if hasattr(source, "read"):
        data = source.read()
    else:
        with open(source, "rb") as f:
            data = f.read()
    return Image.open(io.BytesIO(data)).convert("RGBA")


def paste_overlay(base_img, overlay_img, x_pct, y_pct, size_pct):
    """Paste an RGBA overlay (e.g. a signature) onto a copy of base_img,
    transparency preserved, centered at (x_pct, y_pct) as a percentage of
    the certificate's width/height, sized to size_pct% of the certificate's
    width (aspect ratio kept). Clamped to stay fully within the certificate."""
    base = base_img.convert("RGB").copy()
    overlay = overlay_img.convert("RGBA")

    target_w = max(10, int(base.width * size_pct / 100))
    scale = target_w / overlay.width
    target_h = max(1, int(overlay.height * scale))
    overlay = overlay.resize((target_w, target_h))

    px = int(base.width * x_pct / 100) - overlay.width // 2
    py = int(base.height * y_pct / 100) - overlay.height // 2
    px = max(0, min(px, max(0, base.width - overlay.width)))
    py = max(0, min(py, max(0, base.height - overlay.height)))

    base.paste(overlay, (px, py), overlay)
    return base


# ----------------------------------------------------------------------
# Zoho WorkDrive integration — upload a file and get back a public link,
# so certificate links can be filled in automatically instead of by hand.
# Settings are read per-app via db_path, using the same key/value
# `settings` table as everything else:
#   zoho_client_id, zoho_client_secret, zoho_refresh_token,
#   zoho_accounts_domain (e.g. "accounts.zoho.in"),
#   zoho_workdrive_base   (e.g. "https://workdrive.zoho.in"),
#   zoho_parent_folder_id (WorkDrive folder to upload into)
# ----------------------------------------------------------------------

def zoho_configured(db_path):
    keys = ["zoho_client_id", "zoho_client_secret", "zoho_refresh_token", "zoho_parent_folder_id"]
    return all(get_setting(db_path, k, "").strip() for k in keys)


def zoho_get_access_token(db_path):
    """Exchanges the stored refresh token for a short-lived access token.
    Returns (access_token, error) — error is None on success."""
    client_id = get_setting(db_path, "zoho_client_id", "").strip()
    client_secret = get_setting(db_path, "zoho_client_secret", "").strip()
    refresh_token = get_setting(db_path, "zoho_refresh_token", "").strip()
    accounts_domain = get_setting(db_path, "zoho_accounts_domain", "accounts.zoho.in").strip()
    if not (client_id and client_secret and refresh_token):
        return None, "Zoho credentials are not fully set in Settings."
    try:
        resp = requests.post(
            f"https://{accounts_domain}/oauth/v2/token",
            data={
                "client_id": client_id,
                "client_secret": client_secret,
                "grant_type": "refresh_token",
                "refresh_token": refresh_token,
            },
            timeout=20,
        )
        data = resp.json()
    except Exception as e:
        return None, f"Could not reach Zoho ({accounts_domain}): {e}"
    if "access_token" not in data:
        return None, f"Zoho token refresh failed: {data}"
    return data["access_token"], None


def zoho_exchange_grant_code(client_id, client_secret, grant_code, accounts_domain):
    """One-time exchange of a Self Client grant/authorization code for a
    refresh token. Returns (refresh_token, access_token, error)."""
    try:
        resp = requests.post(
            f"https://{accounts_domain}/oauth/v2/token",
            data={
                "client_id": client_id,
                "client_secret": client_secret,
                "grant_type": "authorization_code",
                "code": grant_code,
            },
            timeout=20,
        )
        data = resp.json()
    except Exception as e:
        return None, None, f"Could not reach Zoho ({accounts_domain}): {e}"
    if "refresh_token" not in data:
        return None, None, f"Zoho code exchange failed: {data}"
    return data["refresh_token"], data.get("access_token"), None


def zoho_upload_file(db_path, file_path, filename=None):
    """Uploads a local file into the configured WorkDrive folder.
    Returns (file_id, error)."""
    token, err = zoho_get_access_token(db_path)
    if err:
        return None, err
    workdrive_base = get_setting(db_path, "zoho_workdrive_base", "https://workdrive.zoho.in").rstrip("/")
    parent_id = get_setting(db_path, "zoho_parent_folder_id", "").strip()
    if not parent_id:
        return None, "Zoho target folder ID is not set in Settings."
    filename = filename or os.path.basename(file_path)
    try:
        with open(file_path, "rb") as f:
            resp = requests.post(
                f"{workdrive_base}/api/v1/upload",
                headers={"Authorization": f"Zoho-oauthtoken {token}"},
                params={"parent_id": parent_id, "filename": filename, "override-name-exist": "true"},
                files={"content": (filename, f)},
                timeout=90,
            )
        data = resp.json()
    except Exception as e:
        return None, f"Upload failed: {e}"

    # Response shape is best-effort from Zoho's docs — fall back through a
    # few plausible paths and surface the raw response if none match, so a
    # mismatch is easy to diagnose rather than failing silently.
    try:
        entry = data["data"][0] if isinstance(data.get("data"), list) else data["data"]
        file_id = entry.get("attributes", {}).get("resource_id") or entry.get("id")
        if not file_id:
            raise KeyError("resource_id/id")
        return file_id, None
    except Exception:
        return None, f"Unexpected upload response — please share this for debugging: {data}"


def zoho_create_public_link(db_path, file_id, link_name):
    """Creates an external 'anyone with the link' share link for a WorkDrive
    file. Returns (url, error). role_id=7 is best-effort for a public,
    no-login-required view link — verify the first real link opens in a
    private/incognito browser window; if it prompts for Zoho login, this
    role_id needs adjusting."""
    token, err = zoho_get_access_token(db_path)
    if err:
        return None, err
    workdrive_base = get_setting(db_path, "zoho_workdrive_base", "https://workdrive.zoho.in").rstrip("/")
    payload = {
        "data": {
            "attributes": {
                "resource_id": file_id,
                "link_name": link_name,
                "request_user_data": False,
                "allow_download": True,
                "role_id": "7",
            },
            "type": "links",
        }
    }
    try:
        resp = requests.post(
            f"{workdrive_base}/api/v1/links",
            headers={
                "Authorization": f"Zoho-oauthtoken {token}",
                "Accept": "application/vnd.api+json",
                "Content-Type": "application/vnd.api+json",
            },
            json=payload,
            timeout=30,
        )
        data = resp.json()
    except Exception as e:
        return None, f"Link creation failed: {e}"

    try:
        attrs = data["data"]["attributes"]
        url = attrs.get("link") or attrs.get("url") or attrs.get("permalink")
        if not url:
            raise KeyError("link/url/permalink")
        return url, None
    except Exception:
        return None, f"Unexpected link response — please share this for debugging: {data}"


def zoho_upload_and_link(db_path, file_path, filename=None):
    """Convenience: upload a file and immediately create its public link.
    Returns (url, error)."""
    filename = filename or os.path.basename(file_path)
    file_id, err = zoho_upload_file(db_path, file_path, filename)
    if err:
        return None, err
    return zoho_create_public_link(db_path, file_id, filename)


def zoho_test_connection(db_path):
    """Read-only check: confirms the token works and the target folder is
    reachable. Returns (ok, message)."""
    token, err = zoho_get_access_token(db_path)
    if err:
        return False, err
    workdrive_base = get_setting(db_path, "zoho_workdrive_base", "https://workdrive.zoho.in").rstrip("/")
    parent_id = get_setting(db_path, "zoho_parent_folder_id", "").strip()
    try:
        resp = requests.get(
            f"{workdrive_base}/api/v1/files/{parent_id}",
            headers={"Authorization": f"Zoho-oauthtoken {token}"},
            timeout=20,
        )
    except Exception as e:
        return False, f"Could not reach Zoho: {e}"
    if resp.status_code == 200:
        return True, "Connected — token and folder ID are both valid."
    return False, f"Zoho responded with status {resp.status_code}: {resp.text[:300]}"
